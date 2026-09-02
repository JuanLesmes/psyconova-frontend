# -*- coding: utf-8 -*-
"""
Convierte la documentación de PSYCONOVA (Markdown) a documentos de Word.

Genera dos archivos:
  - PSYCONOVA - Documentacion Tecnica.docx   (README + capítulos 01 a 12)
  - PSYCONOVA - Propuestas de Mejora.docx    (PROPUESTAS.md)

Los bloques "GRÁFICO G-XX" se convierten en recuadros destacados, para que
sean fáciles de localizar al insertar las imágenes.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ── Paleta de la marca ──────────────────────────────────────────────
TINTA = RGBColor(0x19, 0x20, 0x6D)      # azul tinta, títulos
TEAL = RGBColor(0x0F, 0x8F, 0x84)       # verde azulado, acentos
GRIS = RGBColor(0x5F, 0x67, 0x7E)       # texto secundario
CODIGO = RGBColor(0x8A, 0x2B, 0x4A)     # código en línea
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

SOMBRA_CODIGO = 'F4F6FA'
SOMBRA_GRAFICO = 'FFF8E6'
SOMBRA_TABLA_CAB = '19206D'
SOMBRA_TABLA_ALT = 'F7F8FC'
BORDE_GRAFICO = 'E8B44A'

FUENTE_TEXTO = 'Calibri'
FUENTE_TITULO = 'Calibri Light'
FUENTE_MONO = 'Consolas'

AUTOR = 'Juan Fernando Lesmes Castaneda'
FECHA = '28 de agosto de 2026'


# ── Utilidades de bajo nivel sobre OOXML ────────────────────────────
def sombrear(elemento, hex_color):
    """Aplica color de fondo a una celda o párrafo."""
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    elemento.append(shd)


def sombrear_celda(celda, hex_color):
    sombrear(celda._tc.get_or_add_tcPr(), hex_color)


def sombrear_parrafo(parrafo, hex_color):
    sombrear(parrafo._p.get_or_add_pPr(), hex_color)


def bordes_celda(celda, color, ancho=8, lados=('top', 'left', 'bottom', 'right')):
    tcPr = celda._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for lado in lados:
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(ancho))
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def sin_bordes_tabla(tabla):
    tbl = tabla._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:val'), 'none')
        borders.append(el)
    tblPr.append(borders)


def campo(parrafo, instruccion):
    """Inserta un campo de Word (por ejemplo TOC o PAGE)."""
    r1 = parrafo.add_run()
    fld = OxmlElement('w:fldChar')
    fld.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fld)

    r2 = parrafo.add_run()
    txt = OxmlElement('w:instrText')
    txt.set(qn('xml:space'), 'preserve')
    txt.text = instruccion
    r2._r.append(txt)

    r3 = parrafo.add_run()
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    r3._r.append(sep)

    r4 = parrafo.add_run('Actualiza este campo con F9')
    r4.font.color.rgb = GRIS
    r4.italic = True

    r5 = parrafo.add_run()
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    r5._r.append(end)


def hipervinculo(parrafo, texto, url):
    parte = parrafo.part
    r_id = parte.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True,
    )
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0F8F84')
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(color)
    rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = texto
    run.append(t)
    link.append(run)
    parrafo._p.append(link)


# ── Parseo de Markdown en línea ─────────────────────────────────────
PATRON_INLINE = re.compile(
    r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|`[^`]+`|\[[^\]]*\]\([^)]*\)|\*[^*\s][^*]*?\*)'
)


def agregar_texto(parrafo, texto, base_negrita=False, base_cursiva=False,
                  color=None, tamano=None):
    """Añade texto interpretando negritas, cursivas, código y enlaces."""
    for trozo in PATRON_INLINE.split(texto):
        if not trozo:
            continue

        negrita, cursiva, mono, color_run, contenido = (
            base_negrita, base_cursiva, False, color, trozo
        )

        if trozo.startswith('***') and trozo.endswith('***') and len(trozo) > 6:
            negrita, cursiva, contenido = True, True, trozo[3:-3]
        elif trozo.startswith('**') and trozo.endswith('**') and len(trozo) > 4:
            negrita, contenido = True, trozo[2:-2]
        elif trozo.startswith('`') and trozo.endswith('`') and len(trozo) > 2:
            mono, contenido = True, trozo[1:-1]
            color_run = CODIGO
        elif trozo.startswith('[') and '](' in trozo:
            m = re.match(r'\[([^\]]*)\]\(([^)]*)\)', trozo)
            if m:
                etiqueta, destino = m.group(1), m.group(2)
                if destino.startswith('http'):
                    hipervinculo(parrafo, re.sub(r'[`*]', '', etiqueta), destino)
                    continue
                # Referencia interna: se conserva sólo el texto.
                agregar_texto(parrafo, etiqueta, negrita, True, color, tamano)
                continue
        elif (trozo.startswith('*') and trozo.endswith('*')
              and len(trozo) > 2 and not trozo.startswith('**')):
            cursiva, contenido = True, trozo[1:-1]

        run = parrafo.add_run(contenido)
        run.bold = negrita
        run.italic = cursiva
        if mono:
            run.font.name = FUENTE_MONO
            run.font.size = Pt(9.5)
        if color_run:
            run.font.color.rgb = color_run
        if tamano:
            run.font.size = Pt(tamano)


# ── Construcción del documento ──────────────────────────────────────
class Constructor:
    def __init__(self, titulo_doc):
        self.doc = Document()
        self.titulo_doc = titulo_doc
        self._configurar_estilos()
        self._configurar_pagina()

    def _configurar_estilos(self):
        d = self.doc
        normal = d.styles['Normal']
        normal.font.name = FUENTE_TEXTO
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.12

        encabezados = [
            ('Heading 1', 20, TINTA, 20, 9),
            ('Heading 2', 15, TINTA, 15, 6),
            ('Heading 3', 12.5, TEAL, 12, 4),
            ('Heading 4', 11, GRIS, 10, 3),
        ]
        for nombre, tam, color, antes, despues in encabezados:
            st = d.styles[nombre]
            st.font.name = FUENTE_TITULO if tam >= 15 else FUENTE_TEXTO
            st.font.size = Pt(tam)
            st.font.color.rgb = color
            st.font.bold = True
            st.paragraph_format.space_before = Pt(antes)
            st.paragraph_format.space_after = Pt(despues)
            st.paragraph_format.keep_with_next = True

    def _configurar_pagina(self):
        for s in self.doc.sections:
            s.page_width = Cm(21.59)   # Carta
            s.page_height = Cm(27.94)
            s.top_margin = Cm(2.4)
            s.bottom_margin = Cm(2.2)
            s.left_margin = Cm(2.5)
            s.right_margin = Cm(2.5)

    def propiedades(self, titulo, descripcion):
        cp = self.doc.core_properties
        cp.title = titulo
        cp.author = AUTOR
        cp.comments = descripcion
        cp.category = 'Documentacion'
        cp.language = 'es-CO'

    def pie_y_encabezado(self):
        for s in self.doc.sections:
            s.different_first_page_header_footer = True

            enc = s.header.paragraphs[0]
            enc.text = ''
            enc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = enc.add_run(self.titulo_doc)
            r.font.size = Pt(8)
            r.font.color.rgb = GRIS

            pie = s.footer.paragraphs[0]
            pie.text = ''
            pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = pie.add_run('')
            r.font.size = Pt(9)
            r.font.color.rgb = GRIS
            campo(pie, 'PAGE \\* MERGEFORMAT')
            for run in pie.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = GRIS

    # ── Portada ─────────────────────────────────────────────────
    def portada(self, titulo, subtitulo, lineas_meta):
        d = self.doc
        for _ in range(5):
            d.add_paragraph()

        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run('PSYCONOVA')
        r.font.name = FUENTE_TITULO
        r.font.size = Pt(34)
        r.font.bold = True
        r.font.color.rgb = TINTA

        p = d.add_paragraph()
        r = p.add_run('Psicología · Tecnología · Humanidad')
        r.font.size = Pt(11)
        r.font.color.rgb = TEAL
        r.font.bold = True

        d.add_paragraph()
        linea = d.add_paragraph()
        sombrear_parrafo(linea, '0F8F84')
        linea.paragraph_format.space_after = Pt(2)
        linea.add_run().font.size = Pt(2)

        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(28)
        r = p.add_run(titulo)
        r.font.name = FUENTE_TITULO
        r.font.size = Pt(26)
        r.font.bold = True
        r.font.color.rgb = TINTA

        p = d.add_paragraph()
        r = p.add_run(subtitulo)
        r.font.size = Pt(12.5)
        r.font.color.rgb = GRIS

        for _ in range(6):
            d.add_paragraph()

        for etiqueta, valor in lineas_meta:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f'{etiqueta}   ')
            r.font.size = Pt(9)
            r.font.color.rgb = TEAL
            r.font.bold = True
            r = p.add_run(valor)
            r.font.size = Pt(9.5)
            r.font.color.rgb = GRIS

        d.add_page_break()

    def indice(self, niveles='1-3', salto_final=True):
        p = self.doc.add_paragraph()
        r = p.add_run('Contenido')
        r.font.name = FUENTE_TITULO
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = TINTA
        p.paragraph_format.space_after = Pt(14)

        p = self.doc.add_paragraph()
        campo(p, f'TOC \\o "{niveles}" \\h \\z \\u')

        nota = self.doc.add_paragraph()
        nota.paragraph_format.space_before = Pt(16)
        r = nota.add_run(
            'El índice se rellena solo: haz clic sobre él y pulsa F9, '
            'o Ctrl+E y luego F9 para actualizar todo el documento.'
        )
        r.font.size = Pt(8.5)
        r.font.italic = True
        r.font.color.rgb = GRIS

        if salto_final:
            self.doc.add_page_break()

    # ── Bloques ─────────────────────────────────────────────────
    def bloque_codigo(self, lineas):
        tabla = self.doc.add_table(rows=1, cols=1)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        celda = tabla.cell(0, 0)
        sombrear_celda(celda, SOMBRA_CODIGO)
        bordes_celda(celda, 'DDE3EE', 4)
        celda.text = ''
        for i, linea in enumerate(lineas):
            p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(linea if linea else ' ')
            r.font.name = FUENTE_MONO
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x2A, 0x33, 0x50)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(3)

    @staticmethod
    def agrupar(lineas):
        """
        Une las líneas de continuación de un bloque.

        Hace falta porque una negrita o una cursiva puede partirse en dos
        líneas del Markdown; interpretadas por separado, los asteriscos
        quedarían sueltos y visibles en el documento final.
        """
        grupos = []
        for linea in lineas:
            desnuda = linea.strip()
            if not desnuda:
                continue
            nuevo = (
                desnuda.startswith(('**', '- ', '* ', '#', '|', '>'))
                or re.match(r'^\d+\.\s', desnuda)
                or not grupos
            )
            if nuevo:
                grupos.append(desnuda)
            else:
                grupos[-1] += ' ' + desnuda
        return grupos

    def bloque_grafico(self, lineas):
        """Recuadro destacado para los bloques 'GRÁFICO G-XX'."""
        tabla = self.doc.add_table(rows=1, cols=1)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        celda = tabla.cell(0, 0)
        sombrear_celda(celda, SOMBRA_GRAFICO)
        bordes_celda(celda, BORDE_GRAFICO, 12)
        celda.text = ''

        primera = True
        for linea in self.agrupar(lineas):
            p = celda.paragraphs[0] if primera else celda.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.08

            if primera:
                p.paragraph_format.space_after = Pt(6)
                texto = re.sub(r'^\*\*|\*\*$', '', linea)
                r = p.add_run(texto.replace('📊 ', ''))
                r.font.bold = True
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(0x8D, 0x61, 0x10)
                primera = False
                continue

            m = re.match(r'^[-*]\s+(.*)$', linea)
            if m:
                p.paragraph_format.left_indent = Cm(0.5)
                linea = '•  ' + m.group(1)
            elif re.match(r'^\d+\.\s', linea):
                p.paragraph_format.left_indent = Cm(0.5)

            agregar_texto(p, linea, tamano=9.5)
            for run in p.runs:
                if run.font.size is None:
                    run.font.size = Pt(9.5)

        aviso = celda.add_paragraph()
        aviso.paragraph_format.space_before = Pt(6)
        r = aviso.add_run('▢  Reemplaza este recuadro por la imagen terminada.')
        r.font.size = Pt(8.5)
        r.font.italic = True
        r.font.color.rgb = RGBColor(0xA8, 0x86, 0x40)

        self.doc.add_paragraph().paragraph_format.space_after = Pt(3)

    def bloque_cita(self, lineas):
        for linea in self.agrupar(lineas):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(3)
            pPr = p._p.get_or_add_pPr()
            bordes = OxmlElement('w:pBdr')
            izq = OxmlElement('w:left')
            izq.set(qn('w:val'), 'single')
            izq.set(qn('w:sz'), '18')
            izq.set(qn('w:space'), '10')
            izq.set(qn('w:color'), '0F8F84')
            bordes.append(izq)
            pPr.append(bordes)
            agregar_texto(p, linea, base_cursiva=True, color=GRIS)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def tabla(self, filas):
        cabecera, cuerpo = filas[0], filas[1:]
        t = self.doc.add_table(rows=1, cols=len(cabecera))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, texto in enumerate(cabecera):
            celda = t.rows[0].cells[i]
            celda.text = ''
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            agregar_texto(p, texto, base_negrita=True, color=BLANCO)
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = BLANCO
            sombrear_celda(celda, SOMBRA_TABLA_CAB)

        for n, fila in enumerate(cuerpo):
            celdas = t.add_row().cells
            for i, texto in enumerate(fila[:len(cabecera)]):
                celda = celdas[i]
                celda.text = ''
                p = celda.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                agregar_texto(p, texto)
                for run in p.runs:
                    if run.font.size is None:
                        run.font.size = Pt(9.5)
                if n % 2 == 1:
                    sombrear_celda(celda, SOMBRA_TABLA_ALT)

        self.doc.add_paragraph().paragraph_format.space_after = Pt(3)

    def imagen(self, pie, ruta):
        """
        Inserta un gráfico ya terminado, con su pie.

        Si el archivo todavía no existe se deja un aviso en su lugar, para
        que falte a la vista en vez de desaparecer sin dejar rastro.
        """
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)

        if ruta.exists():
            ancho_util = Cm(21.59 - 2.5 - 2.5)
            p.add_run().add_picture(str(ruta), width=ancho_util)
        else:
            r = p.add_run(f'[falta la imagen: {ruta.name}]')
            r.font.color.rgb = RGBColor(0xC0, 0x3A, 0x3A)
            r.font.italic = True

        if pie:
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(10)
            r = cap.add_run(re.sub(r'[`*]', '', pie))
            r.font.size = Pt(8.5)
            r.font.italic = True
            r.font.color.rgb = GRIS

    def separador(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        pPr = p._p.get_or_add_pPr()
        bordes = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:color'), 'D6DBE8')
        bordes.append(bottom)
        pPr.append(bordes)


# ── Conversión Markdown → documento ─────────────────────────────────
RE_TABLA_SEP = re.compile(r'^\|[\s:|-]+\|$')
RE_ITEM = re.compile(r'^\s*(?:[-*]\s|\d+\.\s)')


def celdas_de(linea):
    return [c.strip() for c in linea.strip().strip('|').split('|')]


def continuacion(lineas, i):
    """
    Recoge las líneas con que sigue un elemento de lista.

    Un elemento de lista puede ocupar varias líneas del Markdown, y una
    negrita partida entre dos de ellas dejaría los asteriscos a la vista si
    cada línea se interpretara por separado.
    """
    extra = []
    while i < len(lineas):
        linea = lineas[i]
        desnuda = linea.strip()
        if (not desnuda or RE_ITEM.match(linea) or not linea[:1].isspace()
                or desnuda.startswith(('#', '>', '|', '```'))):
            break
        extra.append(desnuda)
        i += 1
    return (' ' + ' '.join(extra) if extra else ''), i


def convertir(constructor, texto, base_dir, h1_sub_nivel=2, salto_en_h1=True):
    lineas = texto.split('\n')
    i = 0
    primer_h1 = True

    while i < len(lineas):
        linea = lineas[i]
        desnuda = linea.strip()

        # Bloque de código
        if desnuda.startswith('```'):
            i += 1
            buffer = []
            while i < len(lineas) and not lineas[i].strip().startswith('```'):
                buffer.append(lineas[i])
                i += 1
            i += 1
            constructor.bloque_codigo(buffer)
            continue

        # Cita o bloque de gráfico
        if desnuda.startswith('>'):
            buffer = []
            while i < len(lineas) and lineas[i].strip().startswith('>'):
                buffer.append(re.sub(r'^\s*>\s?', '', lineas[i]))
                i += 1
            es_grafico = any('GRÁFICO G-' in l for l in buffer[:2])
            if es_grafico:
                constructor.bloque_grafico(buffer)
            else:
                constructor.bloque_cita([l for l in buffer if l.strip()])
            continue

        # Tabla
        if (desnuda.startswith('|') and i + 1 < len(lineas)
                and RE_TABLA_SEP.match(lineas[i + 1].strip())):
            filas = [celdas_de(desnuda)]
            i += 2
            while i < len(lineas) and lineas[i].strip().startswith('|'):
                filas.append(celdas_de(lineas[i]))
                i += 1
            constructor.tabla(filas)
            continue

        # Imagen: ![pie](ruta)
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', desnuda)
        if m:
            constructor.imagen(m.group(1), base_dir / m.group(2))
            i += 1
            continue

        # Separador horizontal
        if desnuda in ('---', '***', '___'):
            constructor.separador()
            i += 1
            continue

        # Encabezados
        m = re.match(r'^(#{1,6})\s+(.*)$', desnuda)
        if m:
            nivel_md, texto_h = len(m.group(1)), m.group(2)
            if nivel_md == 1:
                if primer_h1:
                    nivel = 1
                    primer_h1 = False
                    if salto_en_h1 and len(constructor.doc.paragraphs) > 3:
                        constructor.doc.add_page_break()
                else:
                    nivel = h1_sub_nivel
                    if h1_sub_nivel == 1 and salto_en_h1:
                        constructor.doc.add_page_break()
            else:
                nivel = min(nivel_md, 4)

            p = constructor.doc.add_paragraph(style=f'Heading {nivel}')
            agregar_texto(p, texto_h, base_negrita=True)
            estilo = constructor.doc.styles[f'Heading {nivel}']
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = estilo.font.color.rgb
                run.font.size = estilo.font.size
                if run.font.name != FUENTE_MONO:
                    run.font.name = estilo.font.name
            i += 1
            continue

        # Lista con viñetas
        m = re.match(r'^(\s*)[-*]\s+(.*)$', linea)
        if m and not desnuda.startswith('---'):
            sangria = len(m.group(1))
            estilo = 'List Bullet' if sangria < 2 else 'List Bullet 2'
            i += 1
            extra, i = continuacion(lineas, i)
            p = constructor.doc.add_paragraph(style=estilo)
            p.paragraph_format.space_after = Pt(2)
            agregar_texto(p, m.group(2) + extra)
            continue

        # Lista numerada
        m = re.match(r'^(\s*)\d+\.\s+(.*)$', linea)
        if m:
            sangria = len(m.group(1))
            estilo = 'List Number' if sangria < 2 else 'List Number 2'
            i += 1
            extra, i = continuacion(lineas, i)
            p = constructor.doc.add_paragraph(style=estilo)
            p.paragraph_format.space_after = Pt(2)
            agregar_texto(p, m.group(2) + extra)
            continue

        # Línea en blanco
        if not desnuda:
            i += 1
            continue

        # Párrafo (se unen las líneas seguidas)
        buffer = [desnuda]
        i += 1
        while i < len(lineas):
            sig = lineas[i].strip()
            if (not sig or sig.startswith(('#', '>', '|', '```', '- ', '* '))
                    or re.match(r'^\d+\.\s', sig) or sig in ('---', '***', '___')):
                break
            buffer.append(sig)
            i += 1

        p = constructor.doc.add_paragraph()
        agregar_texto(p, ' '.join(buffer))


def limpiar(texto):
    """Quita los punteros de navegación entre capítulos."""
    texto = re.sub(r'\n---\s*\n+\*\*Siguiente:\*\*.*?(?=\n\n|\Z)', '\n',
                   texto, flags=re.S)
    texto = re.sub(r'\n\*\*Siguiente:\*\*.*?(?=\n\n|\Z)', '\n', texto, flags=re.S)
    return texto.rstrip() + '\n'


# ── Programa principal ──────────────────────────────────────────────
def main():
    # Sin argumento, se asume que el script vive en docs/word/ del repositorio.
    raiz = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2]
    docs = raiz / 'docs'
    salida = docs / 'word'
    salida.mkdir(parents=True, exist_ok=True)

    fecha = FECHA

    # ── Documento 1: manual técnico ──────────────────────────────
    manual = Constructor('PSYCONOVA · Documentación técnica')
    manual.portada(
        'Documentación técnica',
        'Sitio web · Manual completo para el equipo de desarrollo',
        [
            ('VERSIÓN', '1.0'),
            ('FECHA', fecha),
            ('REPOSITORIO', 'github.com/Inti-Nova/psyconova-frontend'),
            ('COMMIT', 'bf22a17'),
            ('STACK', 'Angular 21 · TypeScript 5.9 · SCSS · Netlify'),
        ],
    )
    manual.indice('1-3')

    capitulos = [
        docs / 'README.md',
        docs / '01-vision-general.md',
        docs / '02-primeros-pasos.md',
        docs / '03-arquitectura.md',
        docs / '04-sistema-de-diseno.md',
        docs / '05-catalogo-de-componentes.md',
        docs / '06-internacionalizacion.md',
        docs / '07-formulario-de-contacto.md',
        docs / '08-contenido-y-assets.md',
        docs / '09-despliegue-y-operacion.md',
        docs / '10-privacidad-y-legal.md',
        docs / '11-calidad-pruebas-convenciones.md',
        docs / '12-runbook.md',
    ]

    for n, ruta in enumerate(capitulos):
        texto = limpiar(ruta.read_text(encoding='utf-8'))
        if ruta.name == 'README.md':
            texto = texto.replace('# Documentación de PSYCONOVA',
                                  '# Cómo usar esta documentación')
        convertir(manual, texto, docs, h1_sub_nivel=2, salto_en_h1=(n > 0))

    manual.pie_y_encabezado()
    manual.propiedades(
        'PSYCONOVA - Documentacion tecnica',
        'Manual completo del sitio web de PSYCONOVA para el equipo de desarrollo',
    )
    destino = salida / 'PSYCONOVA - Documentacion Tecnica.docx'
    manual.doc.save(str(destino))
    print(f'OK  {destino.name}')

    # ── Documento 2: propuestas ──────────────────────────────────
    prop = Constructor('PSYCONOVA · Propuestas de mejora')
    prop.portada(
        'Propuestas de mejora',
        'Revisión del sitio web y recomendaciones priorizadas',
        [
            ('PREPARADO PARA', 'Laura Valentina Lesmes Castañeda'),
            ('CARGO', 'Fundadora y directora clínica'),
            ('FECHA', fecha),
            ('ALCANCE', 'Revisión completa del sitio publicado'),
        ],
    )
    # Sin salto: cada bloque de nivel 1 ya empieza en página nueva.
    prop.indice('1-2', salto_final=False)

    texto = (docs / 'PROPUESTAS.md').read_text(encoding='utf-8')
    # La portada ya lleva el título y los metadatos: se recorta hasta el
    # primer separador para no repetirlos en la primera página.
    texto = re.sub(r'\A.*?\n---\n', '', texto, count=1, flags=re.S)
    convertir(prop, texto, docs, h1_sub_nivel=1, salto_en_h1=True)

    prop.pie_y_encabezado()
    prop.propiedades(
        'PSYCONOVA - Propuestas de mejora',
        'Revision del sitio web y recomendaciones priorizadas',
    )
    destino = salida / 'PSYCONOVA - Propuestas de Mejora.docx'
    prop.doc.save(str(destino))
    print(f'OK  {destino.name}')


if __name__ == '__main__':
    main()
